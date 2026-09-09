from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("傅思雄_学术实践简历_黑白精简版.docx")
FONT = "SimSun"
BLACK = RGBColor(0, 0, 0)


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def text(p, value, size=9.6, bold=False, italic=False):
    r = p.add_run(value)
    set_font(r, size, bold, italic)
    return r


def set_bottom_rule(p):
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def heading(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    text(p, value, 11.2, True)
    set_bottom_rule(p)


def bullet(doc, value):
    p = doc.add_paragraph(style="Resume Bullet")
    p.paragraph_format.space_after = Pt(1.5)
    text(p, value, 9.4)


def project(doc, name, detail):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.5)
    p.paragraph_format.space_after = Pt(0.5)
    text(p, name, 10, True)
    p = doc.add_paragraph(style="Resume Bullet")
    p.paragraph_format.space_after = Pt(1)
    text(p, detail, 9.4)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.space_after = Pt(2)
    style = doc.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = normal
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.left_indent = Inches(0.24)
    style.paragraph_format.first_line_indent = Inches(-0.17)
    style.paragraph_format.line_spacing = 1.12

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    text(p, "傅思雄", 18, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    text(p, "自动化｜嵌入式系统与机器人控制", 10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    text(p, "北京化工大学 信息学院 ｜ 自动化专业 ｜ 2024级 ｜ 手机：待补充 ｜ 邮箱：待补充", 9)

    heading(doc, "教育背景")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    text(p, "北京化工大学", 10.1, True)
    text(p, " ｜ 信息学院 ｜ 自动化专业 ｜ 2024级 ｜ 自控2401", 10)
    p = doc.add_paragraph()
    text(p, "GPA 3.42 ｜ 大学英语四级 585 ｜ 大学英语六级 504", 9.4)

    heading(doc, "个人简介")
    p = doc.add_paragraph()
    text(p, "注重理论学习与工程实践结合，已围绕 STM32、MSPM0、Arduino/ESP32 平台开展智能车、电机控制、传感器采集和视觉识别实践。希望在嵌入式系统、机器人运动控制、智能感知与机电系统方向继续深入学习。", 9.5)

    heading(doc, "技术能力")
    skills = [
        ("嵌入式开发", "C/C++；STM32F1/F4/G4、TI MSPM0G3507、STC32/8051、Arduino、ESP32；Keil、STM32 HAL、TI DriverLib。"),
        ("外设与通信", "GPIO、ADC、PWM、定时器、DMA、UART、I2C、SPI、CAN；串口协议、遥测帧、蓝牙/Web Serial 联调。"),
        ("控制与机器人", "PID、速度环、航向角/角速度环、级联控制、巡线、编码器测速、IMU 姿态计算；FOC 与 SVPWM 基础。"),
        ("硬件与工具", "PCB 设计、焊接、供电电路、SolidWorks、3D 打印；Python、OpenCV、TensorFlow/Keras、MATLAB/Simulink。"),
    ]
    for label, value in skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        text(p, f"{label}：", 9.4, True)
        text(p, value, 9.4)

    heading(doc, "项目经历")
    project(doc, "MSPM0G3507 智能车控制系统", "完成双电机、编码器、八路灰度、LSM6DSR IMU 与串口上位机集成，实现速度 PID、航向控制、巡线和遥测调试。")
    project(doc, "STM32 智能车与电子技术创新设计", "完成电机、舵机、超声波、OLED 与 PID 等模块开发，参与小车运动控制及软硬件联调。")
    project(doc, "无刷电机 FOC 学习与二次开发", "基于 Arduino/ESP32 学习 AS5600 编码器、速度/位置/电流闭环、Clarke/Park 变换与 SVPWM。")
    project(doc, "TensorFlow 手势识别", "使用 OpenCV、TensorFlow/Keras 和 Tkinter 实现手势数据采集、CNN 训练及实时识别，验证集准确率 99.67%。")
    project(doc, "控制系统建模与仿真", "使用 MATLAB/Simulink 完成 PID、根轨迹和自动驾驶仪控制系统的建模、分析与仿真。")

    heading(doc, "竞赛与荣誉")
    awards = [
        "全国大学生程序设计类竞赛蓝桥杯｜全国一等奖（国一）",
        "全国大学生智能汽车竞赛“飞檐走壁”赛道｜全国三等奖（国三）",
        "全国大学生电子设计竞赛｜省级一等奖（省一）",
        "实用新型专利｜1 项",
    ]
    for award in awards:
        bullet(doc, award)

    heading(doc, "可展示材料")
    p = doc.add_paragraph()
    text(p, "智能车控制代码与调试记录、STM32 小车工程、FOC 学习代码、手势识别运行结果、MATLAB/Simulink 仿真报告及硬件结构设计成果。", 9.4)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    text(p, "注：专利名称、申请号、竞赛年份及联系方式可在提交前补充。", 8.4, italic=True)

    doc.core_properties.title = "傅思雄 学术实践简历"
    doc.core_properties.author = "傅思雄"
    doc.save(OUT)


if __name__ == "__main__":
    build()
