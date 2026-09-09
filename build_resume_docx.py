from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).with_name("傅思雄_学术实践简历.docx")
FONT = "Microsoft YaHei"
ACCENT = "1F4E79"
TEAL = "0F6B6D"
INK = "1D2731"
MUTED = "5D6874"
LIGHT = "EAF1F7"


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell_margins(cell)


def add_text(p, text, size=9.6, bold=False, color=INK, italic=False):
    r = p.add_run(text)
    set_font(r, size, bold, color, italic)
    return r


def add_section(doc, title):
    p = doc.add_paragraph()
    p.style = doc.styles["Resume Heading"]
    add_text(p, title, 12.2, True, ACCENT)
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    p_pr.append(borders)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Resume Bullet")
    p.paragraph_format.space_after = Pt(1.6)
    add_text(p, text, 9.5)
    return p


def add_project(doc, title, role, items):
    p = doc.add_paragraph(style="Project Title")
    add_text(p, title, 10.6, True, INK)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0.5)
    p = doc.add_paragraph(style="Project Meta")
    add_text(p, role, 9.2, False, TEAL)
    p.paragraph_format.space_after = Pt(1)
    for item in items:
        add_bullet(doc, item)


def add_field(paragraph, code):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # Named override: compact academic resume margins, retaining a spacious readable page.
    section.top_margin = Inches(0.56)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, based_on in (("Resume Heading", "Normal"), ("Project Title", "Normal"), ("Project Meta", "Normal"), ("Resume Bullet", "List Bullet")):
        style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[based_on]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Resume Bullet"].paragraph_format.left_indent = Cm(0.46)
    styles["Resume Bullet"].paragraph_format.first_line_indent = Cm(-0.28)
    styles["Resume Bullet"].paragraph_format.line_spacing = 1.12

    # Quiet header/footer for the academic resume.
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header_p, "傅思雄｜学术实践简历", 8.3, False, MUTED)
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer_p, "第 ", 8.2, False, MUTED)
    add_field(footer_p, "PAGE")
    add_text(footer_p, " 页", 8.2, False, MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    add_text(title, "傅思雄", 22, True, ACCENT)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(5)
    add_text(sub, "自动化｜嵌入式系统与机器人控制", 10.5, True, TEAL)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(6)
    add_text(contact, "北京化工大学 信息学院 ｜ 自动化专业 ｜ 2024级 ｜ 手机：待补充 ｜ 邮箱：待补充", 9, False, MUTED)

    add_section(doc, "教育背景")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "北京化工大学", 10.4, True, ACCENT)
    add_text(p, " ｜ 信息学院 ｜ 自动化专业 ｜ 2024级 ｜ 自控2401", 10.0, False)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "GPA 3.42 ｜ 大学英语四级 585 ｜ 大学英语六级 504", 9.4, False, MUTED)

    add_section(doc, "个人简介与学术兴趣")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, "注重理论学习与工程实践结合，已围绕 STM32、MSPM0、Arduino/ESP32 平台完成智能车、电机控制、传感器采集、视觉识别和人机交互项目。具备从电路与 PCB、嵌入式底层驱动、控制算法到上位机联调的完整实践链路，能够进行接口调试、参数整定、故障定位和软硬件协同验证。", 9.6)
    p = doc.add_paragraph()
    add_text(p, "研究兴趣：", 9.6, True, TEAL)
    add_text(p, "嵌入式实时系统、移动机器人与智能车控制、电机驱动与 FOC、传感器融合与姿态估计、机器视觉与智能感知、机电系统软硬件协同设计。", 9.6)

    add_section(doc, "技术能力")
    skills = [
        ("嵌入式", "C/C++；STM32F1/F4/G4、TI MSPM0G3507、STC32/8051、Arduino、ESP32；Keil、STM32 HAL、TI DriverLib、CMSIS。"),
        ("外设与通信", "GPIO、ADC、PWM、定时器、输入捕获、DMA、UART、I2C、SPI、CAN；串口协议、命令解析、遥测帧、蓝牙/Web Serial 联调。"),
        ("控制与机器人", "增量式/位置式 PID、速度环、航向角/角速度环、级联控制、巡线控制、编码器测速、IMU 零偏标定与姿态计算；FOC 电流/速度/位置闭环、SVPWM 基础。"),
        ("电子与机械", "原理图阅读、元器件选型、LM2596S/AMS1117 多级供电、PCB 设计与焊接、硬件故障排查、SolidWorks 建模与 3D 打印。"),
        ("智能算法与工具", "Python、NumPy、OpenCV、TensorFlow/Keras、Tkinter、MATLAB/Simulink、Matplotlib；CNN 图像分类、控制系统仿真与可视化；基础 Git 与 FreeRTOS。"),
    ]
    for label, value in skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        add_text(p, f"{label}：", 9.5, True, TEAL)
        add_text(p, value, 9.5)

    add_section(doc, "项目经历")
    add_project(doc, "MSPM0G3507 智能车控制系统", "核心开发：嵌入式控制、传感器融合、通信与调试", [
        "基于 MSPM0G3507 使用 C 完成双电机智能车控制系统，集成 AT8236 PWM 驱动、双路编码器、八路灰度、LSM6DSR 六轴 IMU、电池 ADC、按键/拨码和蜂鸣器。",
        "实现速度 PI、角速度 PID、航向保持和灰度巡线；采用“航向/角速度 → 左右轮速度 → PWM”级联结构，完成方向切换保护、输出限幅和失联停车。",
        "编写 UART 命令协议与 TEL/STA/SPD/LIN/SQR/DBG 遥测帧，配套 Web Serial/Web Bluetooth 上位机，实现远程控制、状态监控和在线调试。",
        "完成 LSM6DSR SPI 自动探测、256 点陀螺仪零偏标定、静止温漂跟踪、低通滤波及相对航向计算；实测校准编码器参数 1867 counts/m。",
    ])
    add_project(doc, "STM32 智能车与电子技术创新设计", "平台：STM32F103 + HAL｜方向：底层驱动与小车功能集成", [
        "完成电机、舵机、超声波、OLED、滤波和 PID 模块开发；使用 I2C 驱动 SSD1306 OLED，并采用显存缓冲后统一刷新降低闪烁。",
        "参与智能车运动控制、避障/迷宫等功能的软硬件联调，熟悉 GPIO、定时器、PWM、串口和 I2C 外设配置与调试流程。",
    ])
    add_project(doc, "无刷电机 FOC 学习与二次开发", "平台：Arduino/ESP32｜DengFOC", [
        "实践 AS5600 编码器读取、开环/闭环速度、位置闭环、电流力矩环和双路电机控制；理解 Clarke/Park 变换、PID、低通滤波与 SVPWM 的基本流程。",
    ])
    add_project(doc, "基于 TensorFlow 的石头剪刀布手势识别", "技术栈：Python、OpenCV、TensorFlow/Keras、Tkinter", [
        "采集并预处理 rock/paper/scissors 三类手势图像，每类 500 张；从零搭建 CNN 完成训练与实时摄像头识别，验证集准确率 99.67%。",
        "设计倒计时、结果锁定、置信度判断和电脑出拳逻辑，使用 Tkinter 完成交互界面。",
    ])
    add_project(doc, "控制系统建模与仿真", "技术栈：MATLAB、Simulink", [
        "完成零极点、根轨迹和飞机自动驾驶仪 PID 实验，分析控制参数对响应速度、超调量、稳态误差和阻尼的影响，并将仿真结论用于智能车速度/航向环调参。",
    ])

    add_section(doc, "竞赛与荣誉")
    award_rows = [
        ("国家级", "全国大学生程序设计类竞赛蓝桥杯｜全国一等奖（国一）"),
        ("国家级", "全国大学生智能汽车竞赛“飞檐走壁”赛道｜全国三等奖（国三）"),
        ("省级", "全国大学生电子设计竞赛｜省级一等奖（省一）"),
        ("知识产权", "实用新型专利｜1 项"),
    ]
    for level, award in award_rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        add_text(p, f"{level}  ", 9.2, True, TEAL)
        add_text(p, award, 9.5, True)

    add_section(doc, "可展示材料")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "MSPM0G3507 智能车完整说明书、控制代码与上位机调试记录；STM32 小车工程、FOC 学习代码、手势识别运行结果；MATLAB/Simulink 控制仿真报告及电子硬件、结构设计成果。", 9.4)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    add_text(note, "注：专利名称、申请号、竞赛年份及联系方式可在提交前补充。", 8.5, False, MUTED, True)
    doc.core_properties.title = "傅思雄 学术实践简历"
    doc.core_properties.author = "傅思雄"
    doc.save(OUT)


if __name__ == "__main__":
    build()
